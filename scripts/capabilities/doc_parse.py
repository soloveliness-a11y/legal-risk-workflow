"""
文档解析原子能力

职责：
  - 从 PDF / DOCX / TXT / MD 提取纯文本
  - 自动根据扩展名选择解析器
  - 异常处理与降级（PyPDF2 → pypdf）

不依赖：text_filter（纯解析，无筛选逻辑）
"""

import os


def extract_text_from_pdf(file_path):
    """从PDF提取纯文本"""
    try:
        import PyPDF2
        text_parts = []
        with open(file_path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        return '\n\n'.join(text_parts)
    except ImportError:
        print("Warning: PyPDF2 not installed. Trying pypdf...")
        try:
            import pypdf
            text_parts = []
            with open(file_path, 'rb') as f:
                reader = pypdf.PdfReader(f)
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(page_text)
            return '\n\n'.join(text_parts)
        except ImportError:
            print("Error: No PDF library available. Install PyPDF2: pip install PyPDF2")
            return None
    except Exception as e:
        print(f"Error reading PDF: {e}")
        return None


def extract_text_from_docx(file_path):
    """从Word文档提取纯文本（含表格）"""
    try:
        import docx
        document = docx.Document(file_path)
        text_parts = []
        for para in document.paragraphs:
            if para.text.strip():
                text_parts.append(para.text.strip())
        for table in document.tables:
            for row in table.rows:
                row_text = ' | '.join(cell.text.strip() for cell in row.cells)
                if row_text.strip(' |'):
                    text_parts.append(row_text)
        return '\n\n'.join(text_parts)
    except ImportError:
        print("Error: python-docx not installed. Install: pip install python-docx")
        return None
    except Exception as e:
        print(f"Error reading DOCX: {e}")
        return None


def extract_text(file_path):
    """根据文件类型自动选择提取方法"""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == '.pdf':
        return extract_text_from_pdf(file_path)
    elif ext in ('.docx', '.doc'):
        return extract_text_from_docx(file_path)
    elif ext in ('.txt', '.md'):
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
    else:
        print(f"Unsupported file type: {ext}")
        return None
