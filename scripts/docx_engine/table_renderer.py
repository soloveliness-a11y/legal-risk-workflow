"""
表格渲染模块

职责：
  - Markdown 表格解析
  - 列宽自动计算与固定
  - 三种模式的表格渲染（risk_report / enhanced / normal）

依赖：python-docx, constants, oxml_utils, text_utils
"""

from docx.shared import Pt, Twips, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from .constants import RISK_REPORT_STYLE, TABLE_HEADER_BG, TABLE_ALT_ROW_BG, RISK_COLORS, RISK_PREFIX_PATTERN
from .oxml_utils import set_cell_shading, set_cell_margins
from .text_utils import convert_punctuation, parse_inline_formats


def parse_markdown_table(lines):
    """解析 Markdown 表格"""
    rows = []
    for line in lines:
        line = line.strip()
        if line.startswith('|') and line.endswith('|'):
            cells = [c.strip() for c in line.split('|')[1:-1]]
            if all(set(c) <= {'-', ':', ' '} for c in cells):
                continue
            rows.append(cells)
    return rows


def _autofit_then_fix_columns(table, cols, mode="risk_report"):
    """根据内容自动分配列宽，然后锁定为固定列宽。"""
    if mode == "risk_report":
        usable_width_cm = 21.0 - RISK_REPORT_STYLE["margin_left"] - RISK_REPORT_STYLE["margin_right"]
    else:
        usable_width_cm = 21.0 - 3.18 - 3.18

    usable_twips = int(usable_width_cm * 567)
    min_col_twips = int(1.5 * 567)

    col_max_len = [0] * cols
    for row in table.rows:
        for ci in range(min(cols, len(row.cells))):
            text = row.cells[ci].text.strip()
            char_len = sum(2 if ord(c) > 127 else 1 for c in text)
            col_max_len[ci] = max(col_max_len[ci], char_len)

    total_len = max(sum(col_max_len), 1)
    col_widths = []
    for cl in col_max_len:
        w = int(usable_twips * cl / total_len)
        w = max(w, min_col_twips)
        col_widths.append(w)

    total_w = sum(col_widths)
    if total_w > usable_twips:
        scale = usable_twips / total_w
        col_widths = [int(w * scale) for w in col_widths]

    for col_idx, width_twips in enumerate(col_widths):
        if col_idx < len(table.columns):
            for row in table.rows:
                if col_idx < len(row.cells):
                    row.cells[col_idx].width = Twips(width_twips)

    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement('w:tblPr')
    existing = tblPr.find(qn('w:tblLayout'))
    if existing is not None:
        tblPr.remove(existing)
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)


def _set_repeat_header(header_row):
    """设置表头行跨页重复显示（w:tblHeader）。"""
    tr = header_row._tr
    trPr = tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    trPr.append(tblHeader)


def _set_run_font(run, font_en, font_cn, font_size_pt=None):
    """统一设置 run 的字体。font_size_pt 应传裸数值(pt)，如 10.5。"""
    rFonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), font_en)
    rFonts.set(qn('w:hAnsi'), font_cn)
    rFonts.set(qn('w:eastAsia'), font_cn)
    if font_size_pt is not None:
        run.font.size = Pt(font_size_pt)


def _render_cell_content(p, text, font_en, font_cn, font_size_pt, mode="normal"):
    """渲染单元格内容，支持风险标记着色和行内格式。"""
    risk_match = RISK_PREFIX_PATTERN.match(text)
    if risk_match:
        marker = risk_match.group(1)
        color = RISK_COLORS.get(marker)
        rest = text[risk_match.end():]
        if color:
            marker_run = p.add_run(marker)
            marker_run.font.color.rgb = color
            marker_run.bold = True
            _set_run_font(marker_run, font_en, font_cn, font_size_pt)
            segments = parse_inline_formats(rest)
            for seg_text, is_bold, is_italic in segments:
                if not seg_text:
                    continue
                r = p.add_run(convert_punctuation(seg_text))
                r.bold = is_bold
                r.italic = is_italic
                _set_run_font(r, font_en, font_cn, font_size_pt)
            return

    segments = parse_inline_formats(text)
    for seg_text, is_bold, is_italic in segments:
        if not seg_text:
            continue
        r = p.add_run(convert_punctuation(seg_text))
        r.bold = is_bold
        r.italic = is_italic
        _set_run_font(r, font_en, font_cn, font_size_pt)


def add_table(doc, header_row, data_rows, mode="normal"):
    """添加美化表格（表头底纹 + 交替行着色 + 内边距）"""
    cols = len(header_row)
    table = doc.add_table(rows=1 + len(data_rows), cols=cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    if mode == "risk_report":
        tbl_font_size_pt = RISK_REPORT_STYLE["table_font_size_pt"]  # 裸数值，如 10.5
        tbl_line_spacing = RISK_REPORT_STYLE["table_line_spacing"]

        for i, cell_text in enumerate(header_row):
            if i >= cols:
                break
            cell = table.rows[0].cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = tbl_line_spacing
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(convert_punctuation(cell_text.strip()))
            run.bold = True
            _set_run_font(run, RISK_REPORT_STYLE["font_en"], RISK_REPORT_STYLE["font_cn"], tbl_font_size_pt)
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_shading(cell, "D9D9D9")

        for row_idx, row_data in enumerate(data_rows):
            for col_idx, cell_text in enumerate(row_data):
                if col_idx >= cols:
                    break
                cell = table.rows[row_idx + 1].cells[col_idx]
                text = cell_text.strip()
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = tbl_line_spacing
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                _render_cell_content(p, text, RISK_REPORT_STYLE["font_en"],
                                     RISK_REPORT_STYLE["font_cn"], tbl_font_size_pt, mode)
                set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        _autofit_then_fix_columns(table, cols, mode="risk_report")
        _set_repeat_header(table.rows[0])
        return

    # 默认 / enhanced 模式
    for i, cell_text in enumerate(header_row):
        if i >= cols:
            break
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(convert_punctuation(cell_text.strip()))
        run.bold = True
        run.font.size = Pt(10)
        run.font.name = 'Times New Roman'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1E)
        set_cell_shading(cell, TABLE_HEADER_BG)
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    for row_idx, row_data in enumerate(data_rows):
        for col_idx, cell_text in enumerate(row_data):
            if col_idx >= cols:
                break
            cell = table.rows[row_idx + 1].cells[col_idx]
            text = cell_text.strip()
            cell.text = ""
            p = cell.paragraphs[0]
            _render_cell_content(p, text, 'Times New Roman', '宋体', 10, mode)
            if row_idx % 2 == 1:
                set_cell_shading(cell, TABLE_ALT_ROW_BG)
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    _autofit_then_fix_columns(table, cols, mode=mode)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(2)
    spacer.paragraph_format.space_after = Pt(2)
