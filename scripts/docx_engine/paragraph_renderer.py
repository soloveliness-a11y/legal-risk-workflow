"""
段落渲染模块

职责：
  - 正文段落渲染（支持行内格式、风险标记）
  - 列表项渲染（有序/无序，三种模式）
  - run 字体统一设置

依赖：python-docx, constants, oxml_utils, text_utils
"""

from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from .constants import RISK_REPORT_STYLE, RISK_COLORS, RISK_PREFIX_PATTERN
from .oxml_utils import set_first_line_indent_chars, set_hanging_indent_chars, add_list_bullet_style
from .text_utils import convert_punctuation, parse_inline_formats


def set_run_font(run, font_en, font_cn, font_size_pt=None):
    """统一设置 run 的字体（含 hAnsi=中文字体，确保中文引号正确显示）。"""
    rFonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), font_en)
    rFonts.set(qn('w:hAnsi'), font_cn)
    rFonts.set(qn('w:eastAsia'), font_cn)
    if font_size_pt is not None:
        run.font.size = Pt(font_size_pt)


def add_run_with_format(para, text, is_bold=False, is_italic=False,
                        font_cn='宋体', font_en='Times New Roman', font_size=None, mode="normal"):
    """添加带格式的 run，同时自动转换中英文标点。"""
    text = convert_punctuation(text)
    run = para.add_run(text)
    run.bold = is_bold
    run.italic = is_italic

    if mode == "risk_report":
        sz = RISK_REPORT_STYLE["font_size_pt"] if font_size is None else font_size
        set_run_font(run, RISK_REPORT_STYLE["font_en"], RISK_REPORT_STYLE["font_cn"], sz)
    else:
        set_run_font(run, font_en, font_cn, font_size)

    return run


def add_paragraph(doc, text, mode="normal"):
    """添加正文段落，支持行内格式和风险标记"""
    text = text.strip()
    if not text:
        return

    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if mode == "risk_report":
        para.paragraph_format.line_spacing = Pt(28)
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)
        set_first_line_indent_chars(para, 2)

        is_risk_tip = text.startswith('风控提示') or text.startswith('风控建议')
        force_bold = is_risk_tip

        risk_match = RISK_PREFIX_PATTERN.match(text)
        if risk_match:
            marker = risk_match.group(1)
            color = RISK_COLORS.get(marker)
            rest = text[risk_match.end():]
            if color:
                marker_run = add_run_with_format(para, marker, True, False, mode=mode)
                marker_run.font.color.rgb = color
                segments = parse_inline_formats(rest)
                for seg_text, is_bold, is_italic in segments:
                    if seg_text:
                        add_run_with_format(para, seg_text, is_bold or force_bold, is_italic, mode=mode)
                return

        segments = parse_inline_formats(text)
        for seg_text, is_bold, is_italic in segments:
            if not seg_text:
                continue
            add_run_with_format(para, seg_text, is_bold or force_bold, is_italic, mode=mode)
        return

    # 默认 / enhanced 模式
    para.paragraph_format.first_line_indent = Cm(0.74)
    para.paragraph_format.line_spacing = 1.25
    para.paragraph_format.space_before = Pt(3)
    para.paragraph_format.space_after = Pt(3)

    if mode == "enhanced":
        risk_match = RISK_PREFIX_PATTERN.match(text)
        if risk_match:
            marker = risk_match.group(1)
            color = RISK_COLORS.get(marker)
            rest = text[risk_match.end():]
            if color:
                marker_run = add_run_with_format(para, marker, True, False)
                marker_run.font.color.rgb = color
                segments = parse_inline_formats(rest)
                for seg_text, is_bold, is_italic in segments:
                    if seg_text:
                        add_run_with_format(para, seg_text, is_bold, is_italic)
                return

    segments = parse_inline_formats(text)
    for seg_text, is_bold, is_italic in segments:
        if not seg_text:
            continue
        add_run_with_format(para, seg_text, is_bold, is_italic)


def add_list_item(doc, text, ordered=False, number_text=None, mode="normal", is_new_list=False):
    """添加列表项，支持 **加粗** 和 *斜体*。"""
    text = text.strip()
    if not text:
        return

    if mode == "risk_report":
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        para.paragraph_format.line_spacing = Pt(28)
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)

        if ordered:
            set_first_line_indent_chars(para, RISK_REPORT_STYLE["first_line_indent_chars"])
            if number_text:
                num_run = para.add_run(number_text + " ")
                set_run_font(num_run, RISK_REPORT_STYLE["font_en"], RISK_REPORT_STYLE["font_cn"],
                             RISK_REPORT_STYLE["font_size_pt"])
            segments = parse_inline_formats(text)
            for seg_text, is_bold, is_italic in segments:
                if not seg_text:
                    continue
                add_run_with_format(para, seg_text, is_bold, is_italic, mode=mode)
        else:
            set_hanging_indent_chars(para, 2, 1.5)
            add_list_bullet_style(para)
            segments = parse_inline_formats(text)
            for seg_text, is_bold, is_italic in segments:
                if not seg_text:
                    continue
                add_run_with_format(para, seg_text, is_bold, is_italic, mode=mode)
        return

    # 默认模式
    if ordered:
        para = doc.add_paragraph()
        para.paragraph_format.space_before = Pt(1)
        para.paragraph_format.space_after = Pt(1)
        para.paragraph_format.line_spacing = 1.25
        para.paragraph_format.left_indent = Cm(0.74)
        para.paragraph_format.first_line_indent = Cm(-0.74)
        if number_text:
            num_run = para.add_run(number_text + " ")
            set_run_font(num_run, 'Times New Roman', '宋体', 10.5)
        segments = parse_inline_formats(text)
        for seg_text, is_bold, is_italic in segments:
            if not seg_text:
                continue
            add_run_with_format(para, seg_text, is_bold, is_italic)
    else:
        para = doc.add_paragraph(style='List Bullet')
        para.paragraph_format.space_before = Pt(1)
        para.paragraph_format.space_after = Pt(1)
        para.paragraph_format.line_spacing = 1.25
        segments = parse_inline_formats(text)
        for seg_text, is_bold, is_italic in segments:
            if not seg_text:
                continue
            add_run_with_format(para, seg_text, is_bold, is_italic)
